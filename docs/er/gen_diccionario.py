import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from schema import T, ZONAS, REL
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "diccionario_datos_tt2.xlsx")
F = "Arial"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FILL = PatternFill("solid", fgColor="1F3864")
ZFILL = {"transversal": "FFE599", "modulo": "DAE8FC"}
SHEET = {"SEG": "Seguridad", "CAT": "Catalogos", "TRA": "Transversales", "SREA": "SREA",
         "SUG": "SUG", "ENC": "Motor_encuestas", "SES": "SESTAD"}
ORDER = ["SEG", "CAT", "TRA", "SREA", "SUG", "ENC", "SES"]
TS = [("created_at", "TIMESTAMPTZ", "NN", "", "Fecha de creación (Sequelize)."),
      ("updated_at", "TIMESTAMPTZ", "NN", "", "Fecha de última modificación (Sequelize).")]
HEAD = ["Tabla", "#", "Columna", "Tipo PostgreSQL", "PK", "FK", "Único", "Nulo", "Referencia", "Descripción"]
WIDTHS = [24, 5, 28, 16, 5, 5, 8, 6, 26, 70]


def font(**k):
    return Font(name=F, size=k.pop("size", 10), **k)


def header(ws, row, cols):
    for i, h in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER


wb = Workbook()
idx = wb.active
idx.title = "Indice"

# ---------- hojas por zona ----------
for z in ORDER:
    ws = wb.create_sheet(SHEET[z])
    name, kind = ZONAS[z]
    ws["A1"] = f"Diccionario de datos | {name}"
    ws["A1"].font = font(bold=True, size=14)
    ws["A2"] = "P.I.C. TT2 2026-B060 | PostgreSQL 14+ | Fuente: modelo_er_tt2.drawio"
    ws["A2"].font = font(italic=True, color="595959")
    header(ws, 4, HEAD)
    r = 5
    for t, d in T.items():
        if d["zona"] != z:
            continue
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(HEAD))
        c = ws.cell(row=r, column=1, value=f"{t}  |  {d['desc']}")
        c.font = font(bold=True, size=11)
        c.fill = PatternFill("solid", fgColor=ZFILL[kind])
        c.alignment = Alignment(vertical="center")
        ws.row_dimensions[r].height = 20
        r += 1
        cols = d["cols"] + (TS if d["ts"] else [])
        for n, (col, typ, flags, ref, desc) in enumerate(cols, 1):
            f = [x for x in flags.split(",") if x]
            uq = "Sí" if "UQ" in f else ("Comp." if "UQc" in f else "")
            nulo = "No" if ("NN" in f or "PK" in f) else "Sí"
            vals = [t, n, col, typ, "Sí" if "PK" in f else "", "Sí" if "FK" in f else "", uq, nulo, ref, desc]
            for i, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=i, value=v)
                c.font = font(bold=(i == 3 and "PK" in f))
                c.border = BORDER
                c.alignment = Alignment(vertical="top", wrap_text=(i == 10),
                                        horizontal="center" if i in (2, 5, 6, 7, 8) else "left")
            r += 1
        r += 1
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C5"
    ws.sheet_view.showGridLines = False

# ---------- relaciones ----------
wr = wb.create_sheet("Relaciones")
wr["A1"] = "Llaves foráneas del modelo"
wr["A1"].font = font(bold=True, size=14)
rh = ["Tabla hija", "Columna", "Tabla padre", "Columna padre", "Nulo", "Cardinalidad", "Dibujada en el diagrama"]
header(wr, 3, rh)
drawn = {(h, p) for h, p, *_ in REL}
card = {(h, p): (ch, cp) for h, p, ch, cp in REL}
LBL = {"many": "N", "one": "1", "zeroone": "0..1"}
r = 4
for t, d in T.items():
    for col, typ, flags, ref, desc in d["cols"]:
        if not ref:
            continue
        f = flags.split(",")
        pt, pc = ref.split(".")
        nulo = "No" if ("NN" in f or "PK" in f) else "Sí"
        if (t, pt) in card:
            ch, cp = card[(t, pt)]
            cd = f"{LBL[ch]} : {LBL[cp]}"
        else:
            cd = ("0..1 : " if "UQ" in f else "N : ") + ("1" if nulo == "No" else "0..1")
        vals = [t, col, pt, pc, nulo, cd, "Sí" if (t, pt) in drawn else "No (solo atributo)"]
        for i, v in enumerate(vals, 1):
            c = wr.cell(row=r, column=i, value=v)
            c.font = font()
            c.border = BORDER
        r += 1
last_rel = r - 1
for i, w in enumerate([24, 26, 22, 14, 6, 13, 22], 1):
    wr.column_dimensions[get_column_letter(i)].width = w
wr.freeze_panes = "A4"
wr.auto_filter.ref = f"A3:G{last_rel}"
wr.sheet_view.showGridLines = False

# ---------- índice ----------
idx["A1"] = "Diccionario de datos | Plataforma Integrada de Calidad (P.I.C.)"
idx["A1"].font = font(bold=True, size=16)
idx["A2"] = "TT2 2026-B060 | HIMFG | PostgreSQL 14+ | Versión 1.0 | 13/10/2026"
idx["A2"].font = font(italic=True, color="595959")
ih = ["#", "Tabla", "Zona", "Hoja", "Descripción", "Columnas", "FK salientes"]
header(idx, 4, ih)
r = 5
first = r
for i, (t, d) in enumerate(T.items(), 1):
    sh = SHEET[d["zona"]]
    vals = [i, t, ZONAS[d["zona"]][0], sh, d["desc"],
            f"=COUNTIF('{sh}'!$A:$A,B{r})",
            f"=COUNTIF(Relaciones!$A$4:$A${last_rel},B{r})"]
    for j, v in enumerate(vals, 1):
        c = idx.cell(row=r, column=j, value=v)
        c.font = font()
        c.border = BORDER
        c.alignment = Alignment(vertical="top", wrap_text=(j == 5))
    idx.cell(row=r, column=2).hyperlink = f"#'{sh}'!A1"
    idx.cell(row=r, column=2).font = font(color="0563C1", underline="single")
    r += 1
last = r - 1
idx.cell(row=r, column=5, value="Total").font = font(bold=True)
for col in ("F", "G"):
    c = idx[f"{col}{r}"]
    c.value = f"=SUM({col}{first}:{col}{last})"
    c.font = font(bold=True)
    c.border = BORDER
r += 2
notas = [
    "Convenciones",
    "Nulo = 'No' indica restricción NOT NULL. Único = 'Comp.' indica que la columna forma parte de una restricción UNIQUE compuesta.",
    "Las tablas con marca de tiempo incluyen created_at y updated_at (Sequelize, underscored). Las de relación, bitacora, respuesta, concentrado, qsf_seguimiento, encuesta_sestad, encuesta_egreso y folio_secuencia no las incluyen.",
    "Las columnas *_id que apuntan a catalogo_item se validan en la capa de servicios contra el catálogo correspondiente (DEC-11).",
    "Los registros clínicos (incidente, qsf, encuesta) no se eliminan físicamente; los catálogos y usuarios usan desactivación lógica (activo).",
    "Pendientes de validación con el Servicio de Calidad: reinicio anual de folios, catálogos DIMENSION_CALIDAD y CLASIFICACION_QSF, formatos SESTAD de Hospitalización y Urgencias, layouts DGCES de SUG y SESTAD.",
]
for k, n in enumerate(notas):
    c = idx.cell(row=r + k, column=1, value=n)
    c.font = font(bold=(k == 0))
for i, w in enumerate([5, 26, 32, 16, 70, 10, 12], 1):
    idx.column_dimensions[get_column_letter(i)].width = w
idx.freeze_panes = "A5"
idx.sheet_view.showGridLines = False

wb.save(OUT)
print("ok", OUT, "tablas", len(T))
